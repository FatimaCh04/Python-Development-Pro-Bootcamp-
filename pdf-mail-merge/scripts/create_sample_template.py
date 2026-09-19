"""
Create Sample Word Template
Generates templates/employee_template.docx with Jinja2/docxtpl placeholders
for use with the mail merge automation pipeline.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding in twips (1/20th of a point)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex="F2F4F7"):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def create_employee_template(output_path: Path):
    """Programmatically generate a styled Word template containing Jinja2 merge variables."""
    doc = Document()

    # Set page margins (0.75 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Palette
    PRIMARY_COLOR = RGBColor(30, 58, 138)     # Navy Blue
    SECONDARY_COLOR = RGBColor(71, 85, 105)   # Slate Grey
    ACCENT_COLOR = RGBColor(16, 185, 129)     # Emerald Green
    BORDER_COLOR = "CBD5E1"

    # Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("OFFICIAL PERSONNEL RECORD")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("EMPLOYEE INFORMATION & VERIFICATION SUMMARY")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    # Banner / Template Notice Box
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_cell = banner_table.cell(0, 0)
    banner_cell.width = Inches(7.0)
    set_cell_shading(banner_cell, "EFF6FF")  # Light blue background
    set_cell_margins(banner_cell, top=120, bottom=120, left=180, right=180)

    bp = banner_cell.paragraphs[0]
    bp.paragraph_format.space_before = Pt(0)
    bp.paragraph_format.space_after = Pt(0)
    brun1 = bp.add_run("TEMPLATE NOTICE: ")
    brun1.font.name = "Calibri"
    brun1.font.bold = True
    brun1.font.size = Pt(9.5)
    brun1.font.color.rgb = PRIMARY_COLOR
    brun2 = bp.add_run(
        "Fields formatted with double curly braces (e.g. {% raw %}{{ employee_id }}, {{ full_name }}{% endraw %}) "
        "are dynamically populated with verified employee data during the automated mail merge process."
    )
    brun2.font.name = "Calibri"
    brun2.font.italic = True
    brun2.font.size = Pt(9.5)
    brun2.font.color.rgb = SECONDARY_COLOR

    # Metadata Strip (Doc ID, Date, Generated For)
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(14)
    meta_p.paragraph_format.space_after = Pt(10)
    m_run = meta_p.add_run("Document Ref: ")
    m_run.font.bold = True
    m_run.font.size = Pt(9.5)
    m_run2 = meta_p.add_run("EMP-VER-{{ employee_id }}   |   ")
    m_run2.font.size = Pt(9.5)
    m_run3 = meta_p.add_run("Issue Date: ")
    m_run3.font.bold = True
    m_run3.font.size = Pt(9.5)
    m_run4 = meta_p.add_run("{{ current_date }}   |   ")
    m_run4.font.size = Pt(9.5)
    m_run5 = meta_p.add_run("Recipient: ")
    m_run5.font.bold = True
    m_run5.font.size = Pt(9.5)
    m_run6 = meta_p.add_run("{{ full_name }}")
    m_run6.font.size = Pt(9.5)

    def add_section_header(title: str):
        sec_p = doc.add_paragraph()
        sec_p.paragraph_format.space_before = Pt(14)
        sec_p.paragraph_format.space_after = Pt(4)
        run = sec_p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    def create_key_value_table(rows_data):
        table = doc.add_table(rows=len(rows_data), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        col_widths = [Inches(2.5), Inches(4.5)]

        for i, (k, v) in enumerate(rows_data):
            cell_k = table.cell(i, 0)
            cell_v = table.cell(i, 1)

            cell_k.width = col_widths[0]
            cell_v.width = col_widths[1]

            # Shading on alternating rows
            bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
            set_cell_shading(cell_k, bg)
            set_cell_shading(cell_v, bg)

            set_cell_margins(cell_k, top=80, bottom=80, left=140, right=140)
            set_cell_margins(cell_v, top=80, bottom=80, left=140, right=140)

            # Key text
            pk = cell_k.paragraphs[0]
            pk.paragraph_format.space_before = Pt(0)
            pk.paragraph_format.space_after = Pt(0)
            rk = pk.add_run(k)
            rk.font.name = "Calibri"
            rk.font.size = Pt(10)
            rk.font.bold = True
            rk.font.color.rgb = SECONDARY_COLOR

            # Value text (contains docxtpl tag)
            pv = cell_v.paragraphs[0]
            pv.paragraph_format.space_before = Pt(0)
            pv.paragraph_format.space_after = Pt(0)
            rv = pv.add_run(v)
            rv.font.name = "Calibri"
            rv.font.size = Pt(10)
            rv.font.bold = False

        return table

    # 1. Identification
    add_section_header("1. Identification & Contact Information")
    create_key_value_table([
        ("Employee ID Number", "{{ employee_id }}"),
        ("Full Legal Name", "{{ full_name }}"),
        ("Email Address", "{{ email }}"),
        ("Phone Number", "{{ phone }}"),
    ])

    # 2. Position & Organization
    add_section_header("2. Position & Organizational Unit")
    create_key_value_table([
        ("Assigned Department", "{{ department }}"),
        ("Job Designation / Title", "{{ job_title }}"),
        ("Date of Employment / Hire", "{{ hire_date }}"),
        ("Employment Classification", "{{ employment_status }}"),
        ("Direct Reporting Manager", "{{ manager }}"),
    ])

    # 3. Compensation & Performance
    add_section_header("3. Compensation & Review Record")
    create_key_value_table([
        ("Annual Base Salary", "{{ salary }}"),
        ("Performance Tier", "{{ performance_tier }}"),
        ("Work Location / Office", "{{ location }}"),
    ])

    # Statement / Attestation
    add_section_header("4. Administrative Certification")
    cert_p = doc.add_paragraph()
    cert_p.paragraph_format.space_before = Pt(4)
    cert_p.paragraph_format.space_after = Pt(12)
    c_run = cert_p.add_run(
        "This official summary record certifies the current personnel records on file for {{ full_name }} "
        "(Employee ID: {{ employee_id }}) within the {{ department }} department. "
        "Any discrepancies must be reported immediately to the Human Resources Operations Division."
    )
    c_run.font.name = "Calibri"
    c_run.font.size = Pt(9.5)
    c_run.font.color.rgb = SECONDARY_COLOR

    # Signatures Table
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.cell(0, 0).width = Inches(3.5)
    sig_table.cell(0, 1).width = Inches(3.5)
    sig_table.cell(1, 0).width = Inches(3.5)
    sig_table.cell(1, 1).width = Inches(3.5)

    s0 = sig_table.cell(0, 0).paragraphs[0]
    s0.add_run("\n\n_____________________________________\nAuthorized HR Operations Signature")
    s0.runs[0].font.size = Pt(9)
    s0.runs[0].font.color.rgb = SECONDARY_COLOR

    s1 = sig_table.cell(0, 1).paragraphs[0]
    s1.add_run("\n\n_____________________________________\nEmployee Acknowledgment Signature")
    s1.runs[0].font.size = Pt(9)
    s1.runs[0].font.color.rgb = SECONDARY_COLOR

    s2 = sig_table.cell(1, 0).paragraphs[0]
    s2.add_run("Date: {{ current_date }}")
    s2.runs[0].font.size = Pt(9)
    s2.runs[0].font.color.rgb = SECONDARY_COLOR

    s3 = sig_table.cell(1, 1).paragraphs[0]
    s3.add_run("Date: ________________________")
    s3.runs[0].font.size = Pt(9)
    s3.runs[0].font.color.rgb = SECONDARY_COLOR

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"[OK] Successfully generated sample Word template at: {output_path}")


if __name__ == "__main__":
    template_file = Path(__file__).resolve().parent.parent / "templates" / "employee_template.docx"
    create_employee_template(template_file)
